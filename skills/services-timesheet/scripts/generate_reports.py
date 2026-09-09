#!/usr/bin/env python3
"""
generate_reports.py
Generate weekly Excel (column-per-day matrix + audit sections) from services-log.json.

Excel layout:
  Row 1  : Title
  Row 2  : Day headers  (Mon DD/MM ... Fri DD/MM | Total)  -- Notes+Hrs merged per day
  Row 3  : Sub-headers  (Activity/Notes | Hrs per day | Hrs total)
  Rows 4+: Customer blocks, one row per SLOT (slot = position within the busiest day).
           Slot 0 gets Mon event[0], Tue event[0], ... for each day that has an event.
           Slot 1 gets Mon event[1], Tue event[1], ... (or blank if that day has fewer).
           Row count per customer = max events across any single day.
           Customer subtotal row (navy) at bottom of each customer block:
             =SUM formulas per day + week total.
  Final  : TOTAL HOURS row -- =SUM of all customer daily subtotals.
  Below  : Audit sections (Not Matched, Classified Personal).
"""
import argparse
import json
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed.  Run: pip install openpyxl")
    sys.exit(1)

# ---------------------------------------------------------------------------
# Theme
# ---------------------------------------------------------------------------
_C_NAVY  = "1F3864"
_C_PALE  = "EBF3FB"
_C_ALT   = "FFFFFF"
_C_GREY  = "D6DCE4"
_C_WHITE = "FFFFFF"
_C_NM_H  = "ED7D31"
_C_NM_R  = "FFF2CC"
_C_CP_H  = "808080"
_C_CP_R  = "F2F2F2"

_thin   = Side(style="thin", color="AAAAAA")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_CA  = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LA  = Alignment(horizontal="left",   vertical="center", wrap_text=True)
_LA2 = Alignment(horizontal="left",   vertical="top",    wrap_text=True)

def _fill(c):  return PatternFill("solid", fgColor=c)
def _c(ws, r, col, v=None, font=None, fill=None, align=None):
    cell = ws.cell(row=r, column=col, value=v)
    if font:  cell.font      = font
    if fill:  cell.fill      = fill
    if align: cell.alignment = align
    cell.border = _BORDER
    return cell

# ---------------------------------------------------------------------------
# Column layout
# Col  1 = A  : Customer / Cost Object
# Col  2 = B  : Mon Notes     Col  3 = C  : Mon Hrs
# Col  4 = D  : Tue Notes     Col  5 = E  : Tue Hrs
# Col  6 = F  : Wed Notes     Col  7 = G  : Wed Hrs
# Col  8 = H  : Thu Notes     Col  9 = I  : Thu Hrs
# Col 10 = J  : Fri Notes     Col 11 = K  : Fri Hrs
# Col 12 = L  : Week Total Hrs
# ---------------------------------------------------------------------------
TOTAL_COLS = 12

def _notes_col(day_idx): return 2 + day_idx * 2
def _hrs_col(day_idx):   return 3 + day_idx * 2

# ---------------------------------------------------------------------------
# Data helpers
# ---------------------------------------------------------------------------
def load_entries(log: Path, start: date, end: date):
    with open(log, encoding="utf-8") as f:
        all_e = json.load(f)
    out = []
    for e in all_e:
        try:
            d = datetime.strptime(e["date"], "%Y-%m-%d").date()
            if start <= d <= end:
                out.append(e)
        except (KeyError, ValueError):
            continue
    return out


def load_excluded(path: Path, start: date, end: date):
    if not path or not path.exists():
        return []
    with open(path, encoding="utf-8") as f:
        all_excl = json.load(f)
    out = []
    for e in all_excl:
        try:
            d = datetime.strptime(e["date"], "%Y-%m-%d").date()
            if start <= d <= end:
                out.append(e)
        except (KeyError, ValueError):
            continue
    return out


def deduplicate_entries(entries):
    """Remove duplicate log entries that share the same
    (date, customer_name, cost_object, normalised notes, source).

    When duplicates collide the entry with the higher hours value is kept
    (a confirmed entry beats a stub).  Ties are broken by the most recent
    updated_at timestamp; first-seen wins when all fields are equal.

    Returns entries in original order.
    """
    seen = {}          # key -> winning entry
    for e in entries:
        key = (
            e.get("date", ""),
            e.get("customer_name", ""),
            e.get("cost_object", ""),
            (e.get("notes") or "").strip().lower(),
            e.get("source", ""),
        )
        if key not in seen:
            seen[key] = e
        else:
            existing = seen[key]
            new_hrs  = float(e.get("hours") or 0)
            old_hrs  = float(existing.get("hours") or 0)
            if new_hrs > old_hrs:
                seen[key] = e
            elif new_hrs == old_hrs and (
                e.get("updated_at", "") > existing.get("updated_at", "")
            ):
                seen[key] = e
    return list(seen.values())


def group_entries(entries):
    """Group by (customer, cost_object, cost_object_type); sort within group by date."""
    g = defaultdict(list)
    for e in entries:
        key = (
            e.get("customer_name", "Unknown"),
            e.get("cost_object", ""),
            e.get("cost_object_type", ""),
        )
        g[key].append(e)
    for key in g:
        g[key].sort(key=lambda x: (
            x.get("date", ""),
            0 if x.get("source") == "calendar" else 1,
            x.get("notes") or "",
        ))
    return g


# ---------------------------------------------------------------------------
# Audit section
# ---------------------------------------------------------------------------
def write_excluded_section(ws, start_row, items, bucket):
    if not items:
        return start_row
    row    = start_row
    F_HEAD = Font(bold=True, color=_C_WHITE, size=10)
    F_SUBH = Font(bold=True, size=9)
    F_ROW  = Font(size=9, italic=True, color="595959")

    if bucket == "not_matched":
        h_col = _C_NM_H;  r_col = _C_NM_R
        title = f"  NOT MATCHED  -  {len(items)} item(s)  -  shown for transparency, not logged"
    else:
        h_col = _C_CP_H;  r_col = _C_CP_R
        title = f"  CLASSIFIED PERSONAL  -  {len(items)} item(s)  -  shown for transparency, not logged"

    ws.row_dimensions[row].height = 6
    row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=TOTAL_COLS)
    h = ws.cell(row=row, column=1, value=title)
    h.font = F_HEAD;  h.fill = _fill(h_col);  h.alignment = _LA;  h.border = _BORDER
    ws.row_dimensions[row].height = 18;  row += 1

    for col, hdr in [(1, "Date"), (2, "Activity / Notes"), (3, "Src"), (4, "Reason")]:
        c = ws.cell(row=row, column=col, value=hdr)
        c.font = F_SUBH;  c.fill = _fill(r_col);  c.alignment = _CA;  c.border = _BORDER
    ws.row_dimensions[row].height = 14;  row += 1

    alt       = _fill(r_col)
    last_date = None
    for item in sorted(items, key=lambda x: (x.get("date", ""), x.get("source", ""), x.get("notes", ""))):
        d_str = item.get("date", "")
        if d_str != last_date:
            try:    dl = datetime.strptime(d_str, "%Y-%m-%d").strftime("%a %d/%m")
            except: dl = d_str
            last_date = d_str
        else:
            dl = ""
        src    = "Cal" if item.get("source") == "calendar" else "Email"
        notes  = (item.get("notes")  or "").strip()[:100]
        reason = (item.get("reason") or "").strip()[:120]
        _c(ws, row, 1, dl or None,     F_ROW, alt, _CA)
        _c(ws, row, 2, notes or None,  F_ROW, alt, _LA2)
        _c(ws, row, 3, src,            F_ROW, alt, _CA)
        _c(ws, row, 4, reason or None, F_ROW, alt, _LA2)
        ws.row_dimensions[row].height = 15;  row += 1

    return row


# ---------------------------------------------------------------------------
# Excel writer  (column-per-day matrix, slot-based rows)
# ---------------------------------------------------------------------------
def write_excel(entries, week_dates, out_path, excluded=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Timesheet"

    F_TITLE = Font(bold=True, size=13, color=_C_NAVY)
    F_DAYH  = Font(bold=True, color=_C_WHITE, size=9)
    F_SUBH  = Font(bold=True, size=8, color=_C_NAVY)
    F_NORM  = Font(size=9)
    F_STUB  = Font(size=9, italic=True, color="595959")
    F_TOT   = Font(bold=True, color=_C_WHITE, size=9)
    F_GRAND = Font(bold=True, color=_C_WHITE, size=10)

    s, e_d     = week_dates[0], week_dates[-1]
    wlabel     = f"{s.strftime('%d %b %Y')} to {e_d.strftime('%d %b %Y')}"
    day_labels = [d.strftime("%a  %d/%m") for d in week_dates]
    day_iso    = [d.isoformat() for d in week_dates]

    # -- Row 1 : Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=TOTAL_COLS)
    t = ws.cell(row=1, column=1, value=f"Weekly Services Timesheet  -  {wlabel}")
    t.font = F_TITLE;  t.alignment = _CA
    ws.row_dimensions[1].height = 26

    # -- Row 2 : Day headers
    _c(ws, 2, 1, "Customer / Cost Object", F_DAYH, _fill(_C_NAVY), _CA)
    for i, dlbl in enumerate(day_labels):
        nc = _notes_col(i);  hc = _hrs_col(i)
        ws.merge_cells(start_row=2, start_column=nc, end_row=2, end_column=hc)
        _c(ws, 2, nc, dlbl, F_DAYH, _fill(_C_NAVY), _CA)
    _c(ws, 2, TOTAL_COLS, "Total", F_DAYH, _fill(_C_NAVY), _CA)
    ws.row_dimensions[2].height = 20

    # -- Row 3 : Sub-headers
    _c(ws, 3, 1, "", F_SUBH, _fill(_C_GREY), _CA)
    for i in range(5):
        nc = _notes_col(i);  hc = _hrs_col(i)
        _c(ws, 3, nc, "Activity / Notes", F_SUBH, _fill(_C_GREY), _LA)
        _c(ws, 3, hc, "Hrs",              F_SUBH, _fill(_C_GREY), _CA)
    _c(ws, 3, TOTAL_COLS, "Hrs", F_SUBH, _fill(_C_GREY), _CA)
    ws.row_dimensions[3].height = 15

    groups            = group_entries(entries)
    row               = 4
    cust_subtotal_rows = []

    # -- Customer blocks
    for (customer, cost_obj, obj_type) in sorted(groups.keys()):
        evts     = groups[(customer, cost_obj, obj_type)]
        cust_lbl = f"{customer}\n[{cost_obj}]" if cost_obj else customer

        day_evts = defaultdict(list)
        for ent in evts:
            d_str = ent.get("date", "")
            if d_str in day_iso:
                day_evts[day_iso.index(d_str)].append(ent)

        max_slots = max((len(day_evts[i]) for i in range(5)), default=0)
        max_slots = max(max_slots, 1)

        evt_first_row = row

        for slot in range(max_slots):
            slot_fill = _fill(_C_PALE) if slot % 2 == 0 else _fill(_C_ALT)
            _c(ws, row, 1, None, F_NORM, slot_fill, _LA)

            for day_idx in range(5):
                nc  = _notes_col(day_idx)
                hc  = _hrs_col(day_idx)
                ent = day_evts[day_idx][slot] if slot < len(day_evts[day_idx]) else None

                if ent is None:
                    _c(ws, row, nc, None, F_NORM, slot_fill, _LA2)
                    _c(ws, row, hc, None, F_NORM, slot_fill, _CA)
                else:
                    hrs     = float(ent.get("hours") or 0)
                    notes   = (ent.get("notes") or "").strip()
                    is_stub = hrs == 0.0
                    src_ico = "Cal" if ent.get("source") == "calendar" else "Email"
                    label   = f"{src_ico} {notes}" if notes else src_ico
                    if is_stub:
                        label += "  > fill hrs"
                    f = F_STUB if is_stub else F_NORM
                    _c(ws, row, nc, label or None,            f, slot_fill, _LA2)
                    _c(ws, row, hc, hrs if hrs > 0 else None, f, slot_fill, _CA)

            _c(ws, row, TOTAL_COLS, None, F_NORM, slot_fill, _CA)
            ws.row_dimensions[row].height = 20
            row += 1

        evt_last_row = row - 1

        _c(ws, row, 1, cust_lbl, F_TOT, _fill(_C_NAVY), _LA)
        day_subtotal_cells = {}
        for day_idx in range(5):
            nc     = _notes_col(day_idx)
            hc     = _hrs_col(day_idx)
            hc_ltr = get_column_letter(hc)
            _c(ws, row, nc, None, F_TOT, _fill(_C_NAVY), _CA)
            formula = f"=SUM({hc_ltr}{evt_first_row}:{hc_ltr}{evt_last_row})"
            _c(ws, row, hc, formula, F_TOT, _fill(_C_NAVY), _CA)
            day_subtotal_cells[day_idx] = f"{hc_ltr}{row}"
        week_refs = ",".join(day_subtotal_cells[i] for i in range(5))
        _c(ws, row, TOTAL_COLS, f"=SUM({week_refs})", F_TOT, _fill(_C_NAVY), _CA)
        ws.row_dimensions[row].height = 22
        cust_subtotal_rows.append((row, day_subtotal_cells))
        row += 1

        for col in range(1, TOTAL_COLS + 1):
            ws.cell(row=row, column=col).border = _BORDER
        ws.row_dimensions[row].height = 4
        row += 1

    # -- TOTAL HOURS row
    _c(ws, row, 1, "TOTAL HOURS", F_GRAND, _fill(_C_NAVY), _CA)
    for day_idx in range(5):
        nc     = _notes_col(day_idx)
        hc     = _hrs_col(day_idx)
        hc_ltr = get_column_letter(hc)
        _c(ws, row, nc, None, F_GRAND, _fill(_C_NAVY), _CA)
        day_refs = ",".join(d[1][day_idx] for d in cust_subtotal_rows if day_idx in d[1])
        _c(ws, row, hc, f"=SUM({day_refs})" if day_refs else 0, F_GRAND, _fill(_C_NAVY), _CA)
    tot_refs = ",".join(get_column_letter(TOTAL_COLS) + str(d[0]) for d in cust_subtotal_rows)
    _c(ws, row, TOTAL_COLS, f"=SUM({tot_refs})" if tot_refs else 0, F_GRAND, _fill(_C_NAVY), _CA)
    ws.row_dimensions[row].height = 22
    row += 1

    # -- Audit sections
    if excluded:
        nm  = [e for e in excluded if e.get("bucket") == "not_matched"]
        cp  = [e for e in excluded if e.get("bucket") == "classified_personal"]
        row = write_excluded_section(ws, row, nm, "not_matched")
        row = write_excluded_section(ws, row, cp, "classified_personal")

    # -- Column widths & freeze
    ws.column_dimensions["A"].width = 22
    for i in range(5):
        ws.column_dimensions[get_column_letter(_notes_col(i))].width = 42
        ws.column_dimensions[get_column_letter(_hrs_col(i))].width   = 6
    ws.column_dimensions[get_column_letter(TOTAL_COLS)].width = 7
    ws.freeze_panes = ws.cell(row=4, column=2)

    wb.save(out_path)
    stubs   = sum(1 for e in entries if float(e.get("hours") or 0) == 0)
    total   = sum(float(e.get("hours") or 0) for e in entries)
    excl_nm = len([e for e in (excluded or []) if e.get("bucket") == "not_matched"])
    excl_cp = len([e for e in (excluded or []) if e.get("bucket") == "classified_personal"])
    print(f"OK: Excel saved  -> {out_path}")
    print(f"    Logged: {len(entries)} entries | {total:.1f}h | {stubs} stubs")
    if excluded:
        print(f"    Audit:  {excl_nm} not-matched | {excl_cp} classified personal")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    p = argparse.ArgumentParser(description="Generate services timesheet Excel report")
    p.add_argument("--log",      required=True)
    p.add_argument("--start",    required=True)
    p.add_argument("--end",      required=True)
    p.add_argument("--excel",    required=True)
    p.add_argument("--excluded", default=None)
    args = p.parse_args()

    log_path = Path(args.log)
    if not log_path.exists():
        print(f"ERROR: {log_path} not found");  sys.exit(1)

    try:
        start = datetime.strptime(args.start, "%Y-%m-%d").date()
        end   = datetime.strptime(args.end,   "%Y-%m-%d").date()
    except ValueError as ex:
        print(f"ERROR: {ex}");  sys.exit(1)

    week_dates = [start + timedelta(days=i) for i in range(5)]
    raw        = load_entries(log_path, start, end)
    entries    = deduplicate_entries(raw)
    excluded   = load_excluded(Path(args.excluded), start, end) if args.excluded else []

    dup_count = len(raw) - len(entries)
    if dup_count:
        print(f"  Note: {dup_count} duplicate entr{'y' if dup_count == 1 else 'ies'} removed (same date/customer/notes/source)")

    if not entries:
        print(f"No entries for {args.start} to {args.end}");  sys.exit(0)

    write_excel(entries, week_dates, Path(args.excel), excluded)


if __name__ == "__main__":
    main()
